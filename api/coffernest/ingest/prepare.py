"""Prepare a dropped file for the model: text where we can get it, page images where we can't.

Text is cheaper and more exact than vision, so a statement PDF goes in as text; scanned pages,
photos and screenshots go in as JPEGs no larger than the model would downscale to anyway.
PDFium is not thread-safe, so every PDF call runs under one lock inside a worker thread.
"""

from __future__ import annotations

import csv
import io
import threading
from collections.abc import Sequence
from dataclasses import dataclass, field

MAX_IMAGES = 6
MAX_EDGE = 1568
MAX_TEXT = 120_000
MAX_ROWS = 3000
_PDF_LOCK = threading.Lock()

IMAGE_MIMES = {"image/jpeg", "image/png", "image/webp", "image/heic", "image/heif", "image/gif"}
XLSX_MIMES = {
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "application/vnd.ms-excel",
}


class Locked(Exception):
    """The PDF needs a password none of the saved ones opened."""


class Unreadable(Exception):
    """The file is damaged or of a type we can't read."""


@dataclass
class Prepared:
    kind: str  # pdf | image | sheet | text
    text: str = ""
    images: list[bytes] = field(default_factory=list)  # JPEG bytes
    pages: int = 0
    password_index: int | None = None  # which saved password opened it


def sniff(filename: str, mime: str, data: bytes) -> str:
    name = filename.lower()
    head = data[:8]
    if head.startswith(b"%PDF") or name.endswith(".pdf"):
        return "pdf"
    if head.startswith(b"PK") and name.endswith((".xlsx", ".xlsm")):
        return "xlsx"
    if mime in XLSX_MIMES or name.endswith((".xlsx", ".xlsm")):
        return "xlsx"
    if (
        mime in IMAGE_MIMES
        or head.startswith((b"\xff\xd8", b"\x89PNG", b"GIF8"))
        or (head[:4] == b"RIFF" and data[8:12] == b"WEBP")
        or name.endswith((".jpg", ".jpeg", ".png", ".webp", ".heic", ".heif"))
    ):
        return "image"
    if name.endswith((".csv", ".tsv", ".txt")) or mime.startswith("text/"):
        return "text"
    return "unknown"


def _jpeg(image: object) -> bytes:
    from PIL import Image

    assert isinstance(image, Image.Image)
    img = image.convert("RGB")
    img.thumbnail((MAX_EDGE, MAX_EDGE))
    out = io.BytesIO()
    img.save(out, format="JPEG", quality=82, optimize=True)
    return out.getvalue()


def _image(data: bytes) -> Prepared:
    from PIL import Image, UnidentifiedImageError

    try:
        import pillow_heif

        pillow_heif.register_heif_opener()
    except ImportError:  # HEIC just won't open; browsers usually convert anyway
        pass
    try:
        with Image.open(io.BytesIO(data)) as img:
            img.seek(0)
            return Prepared(kind="image", images=[_jpeg(img)], pages=1)
    except (UnidentifiedImageError, OSError) as exc:
        raise Unreadable("That image couldn't be opened") from exc


def _pdf(data: bytes, passwords: Sequence[str]) -> Prepared:
    import pypdfium2 as pdfium

    with _PDF_LOCK:
        doc = None
        used: int | None = None
        for index, password in enumerate([None, *passwords]):
            try:
                doc = pdfium.PdfDocument(data, password=password)
                used = None if password is None else index - 1
                break
            except pdfium.PdfiumError as exc:
                # FPDF_ERR_PASSWORD == 4; anything else means the file itself is bad.
                if getattr(exc, "err_code", None) not in (4, None) and password is None:
                    raise Unreadable("That PDF is damaged or not really a PDF") from exc
                continue
        if doc is None:
            raise Locked()
        try:
            texts: list[str] = []
            images: list[bytes] = []
            pages = len(doc)
            for i in range(pages):
                page = doc[i]
                try:
                    textpage = page.get_textpage()
                    try:
                        text = textpage.get_text_range().strip()
                    finally:
                        textpage.close()
                    if len(text) >= 80:
                        texts.append(f"--- page {i + 1} ---\n{text}")
                    elif len(images) < MAX_IMAGES:  # scanned page: send a picture of it
                        bitmap = page.render(scale=150 / 72)
                        try:
                            images.append(_jpeg(bitmap.to_pil()))
                        finally:
                            bitmap.close()
                finally:
                    page.close()
                if sum(len(t) for t in texts) > MAX_TEXT:
                    texts.append("--- (later pages cut) ---")
                    break
        finally:
            doc.close()
    return Prepared(
        kind="pdf",
        text="\n".join(texts)[:MAX_TEXT],
        images=images,
        pages=pages,
        password_index=used,
    )


def _rows_to_tsv(rows: list[list[str]]) -> str:
    # Drop leading junk (bank name, address) by starting at the first row that looks like a header.
    keys = (
        "date",
        "narration",
        "description",
        "particulars",
        "debit",
        "credit",
        "amount",
        "withdrawal",
        "deposit",
        "balance",
    )
    start = 0
    for i, row in enumerate(rows[:40]):
        joined = " ".join(row).lower()
        if sum(k in joined for k in keys) >= 2:
            start = i
            break
    kept = [r for r in rows[start:] if any(c.strip() for c in r)][:MAX_ROWS]
    return "\n".join("\t".join(c.strip().replace("\t", " ") for c in r) for r in kept)


def _text(data: bytes) -> Prepared:
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            raw = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    # Bank exports start with name/address lines, which confuse csv.Sniffer; pick the
    # delimiter that splits the most lines into 3+ columns instead.
    lines = raw.splitlines()[:80]
    scores = {d: sum(line.count(d) >= 2 for line in lines) for d in (",", ";", "\t", "|")}
    delimiter, score = max(scores.items(), key=lambda kv: kv[1])
    if score < 2:
        return Prepared(kind="text", text=raw[:MAX_TEXT], pages=1)
    rows = [list(r) for r in csv.reader(io.StringIO(raw), delimiter=delimiter)]
    return Prepared(kind="sheet", text=_rows_to_tsv(rows)[:MAX_TEXT], pages=1)


def _xlsx(data: bytes) -> Prepared:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise Unreadable("That spreadsheet couldn't be opened (is it password-protected?)") from exc
    try:
        parts: list[str] = []
        for ws in wb.worksheets[:5]:
            rows = [
                ["" if v is None else str(v) for v in row]
                for _, row in zip(range(MAX_ROWS), ws.iter_rows(values_only=True), strict=False)
            ]
            if rows:
                parts.append(f"--- sheet {ws.title} ---\n{_rows_to_tsv(rows)}")
    finally:
        wb.close()
    return Prepared(kind="sheet", text="\n".join(parts)[:MAX_TEXT], pages=len(parts))


def prepare(filename: str, mime: str, data: bytes, passwords: Sequence[str] = ()) -> Prepared:
    kind = sniff(filename, mime, data)
    if kind == "pdf":
        return _pdf(data, passwords)
    if kind == "image":
        return _image(data)
    if kind == "xlsx":
        return _xlsx(data)
    if kind == "text":
        return _text(data)
    raise Unreadable("Drop a PDF, a photo or screenshot, a CSV or an Excel file")
